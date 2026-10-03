"""Build small .xlsx files that let the template's formulas be checked in a
REAL Google Sheet, without an API key.

Why this exists: every Sheets bug in this template (serial dates, blank GDP
sectors, a freshness cell stuck on "—", #ERROR! from open ranges) passed the
string tests in tests/sheets/. Only a real Sheet evaluates the formulas.

The raw tab is pre-filled exactly as IMPORTDATA(..., "en_US") leaves it — ISO
dates and YYYY-MM as bare serial numbers, numbers as numbers, text as text —
so everything downstream of the one network call runs for real.

  v1.xlsx  raw tab + Vàng SJC / CPI / GDP / Lãi suất LNH tabs, real formulas
  v2.xlsx  the dashboard over static dataset tabs

Each is kept small (styles and charts stripped, theme dropped): files above
~14k base64 characters failed to upload through the Drive connector.

Usage:
  python3 integrations/google-sheets/make_verify_files.py stacked.csv OUT_DIR
where stacked.csv is the body of /api/v1/excel/refresh-data?format=csv.
Upload each file to Drive converted to a Google Sheet, then read it back —
Drive evaluates the formulas on conversion. Check by eye: dates are text,
GDP "Khu vực" is filled, B2 shows a date, dashboard values match the tabs.
SPARKLINE renders as an image, so a text read only proves it did not error.
"""
import csv
import re
import sys
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_template as bt  # noqa: E402

EPOCH = date(1899, 12, 30)


def as_importdata(v: str):
    if v == "":
        return None
    try:
        if len(v) == 10 and v[4] == "-":
            return (date.fromisoformat(v) - EPOCH).days
        if len(v) == 7 and v[4] == "-":
            return (date(int(v[:4]), int(v[5:]), 1) - EPOCH).days
    except ValueError:
        pass
    try:
        return float(v)
    except ValueError:
        return v


def plain(ws):
    for row in ws.iter_rows():
        for c in row:
            c.style = "Normal"


def slim(src: Path, dst: Path) -> None:
    """Re-zip at max compression without theme/docProps."""
    drop = {"xl/theme/theme1.xml", "docProps/app.xml", "docProps/core.xml"}
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zout:
        for item in zin.infolist():
            if item.filename in drop:
                continue
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = re.sub(rb'<Override PartName="/(xl/theme/theme1|docProps/app|docProps/core)\.xml"[^>]*/>', b"", data)
            elif item.filename == "_rels/.rels":
                data = re.sub(rb'<Relationship [^>]*Target="/?docProps/[^"]*"[^>]*/>', b"", data)
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = re.sub(rb'<Relationship [^>]*Target="/?(xl/)?theme/theme1\.xml"[^>]*/>', b"", data)
            zout.writestr(item.filename, data)


def main(csv_path: str, out_dir: str) -> None:
    out = Path(out_dir)
    rows = list(csv.reader(open(csv_path, encoding="utf-8")))
    by = {}
    for row in rows[1:]:
        by.setdefault(row[0], []).append(row)
    keep = [rows[0]] + [r for k, v in by.items() for r in (v[-5:] if k == "gdp" else v[-3:])]
    bt.MAX_ROWS = 30
    for d in bt.DATASETS:
        d["cols"] = [(h, c, k, None) for h, c, k, f in d["cols"]]

    def add_names(wb):
        from openpyxl.workbook.defined_name import DefinedName
        for name, ref in bt.raw_names().items():
            wb.defined_names[name] = DefinedName(name, attr_text=ref)

    wb = Workbook()
    wb.remove(wb.active)
    for tid in ["gold", "cpi", "gdp", "interbank"]:
        ws = wb.create_sheet(bt.BY_ID[tid]["tab"])
        bt.build_dataset_tab(ws, bt.BY_ID[tid])
        if tid == "gdp":
            bt.build_gdp_helper(ws)
        plain(ws)
    raw = wb.create_sheet(bt.RAW)
    for r, row in enumerate(keep, start=1):
        for c, v in enumerate(row, start=1):
            raw.cell(row=r, column=c, value=v if r == 1 else as_importdata(v))
    add_names(wb)
    wb.save(out / "v1_full.xlsx")
    slim(out / "v1_full.xlsx", out / "v1.xlsx")

    small = {}
    for row in keep[1:]:
        small.setdefault(row[0], []).append(row)
    wb = Workbook()
    wb.remove(wb.active)
    dash = wb.create_sheet(bt.DASH)
    tabs = {}
    for spec in bt.DATASETS:
        ws = wb.create_sheet(spec["tab"])
        tabs[spec["id"]] = ws
        ws["B2"] = small[spec["id"]][-1][1]
        for i, row in enumerate(small[spec["id"]]):
            for j, (_h, col, kind, _f) in enumerate(spec["cols"], start=1):
                v = row[ord(col) - ord("A")]
                ws.cell(row=bt.DATA_ROW + i, column=j,
                        value=v if kind in ("date", "month", "sector") or v == "" else float(v))
    tabs["gdp"]["I6"] = 7.83
    raw = wb.create_sheet(bt.RAW)
    raw["A1"] = "dataset"
    bt.build_dashboard(dash, tabs)
    plain(dash)
    add_names(wb)
    wb.save(out / "v2_full.xlsx")
    slim(out / "v2_full.xlsx", out / "v2.xlsx")
    print(out / "v1.xlsx", out / "v2.xlsx")


if __name__ == "__main__":
    main(*sys.argv[1:3])
