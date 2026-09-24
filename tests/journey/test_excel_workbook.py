"""Workbook refresh contract without requiring Excel, VBA or a live database."""
import asyncio
import io
import sys
import unittest
import zipfile
from datetime import date
from pathlib import Path
from unittest.mock import AsyncMock, patch

from fastapi.encoders import jsonable_encoder
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "be"))

from routers import excel_export


class _Request:
    base_url = "https://api.example/"


def _datasets():
    items = []
    for dataset_id, sheet, table, _fn, _kwargs, description, endpoint in excel_export._datasets("1y"):
        headers = ["period", "value"] if dataset_id == "cpi" else ["date", "value"]
        first = "2026-09" if dataset_id == "cpi" else date(2026, 9, 23)
        items.append({
            "id": dataset_id,
            "sheet": sheet,
            "table": table,
            "description": description,
            "endpoint": f"/api/v1/{endpoint}",
            "headers": headers,
            "rows": [[first, 123.5]],
            "error": None,
        })
    return items


class ExcelWorkbookTests(unittest.TestCase):
    def _response(self):
        data = _datasets()
        with patch.object(excel_export, "_collect_datasets", AsyncMock(return_value=data)):
            return asyncio.run(excel_export.excel_workbook(_Request(), "1y"))

    def test_workbook_has_refreshable_tables_and_real_types(self):
        response = self._response()
        wb = load_workbook(io.BytesIO(response.body))

        self.assertEqual(wb.sheetnames[0], "Bắt đầu")
        self.assertEqual(len(wb.sheetnames), 10)
        self.assertIsInstance(wb["Vàng SJC"]["A5"].value, date)
        self.assertEqual(wb["Vàng SJC"]["A5"].value.date(), date(2026, 9, 23))
        self.assertIsInstance(wb["Vàng SJC"]["B5"].value, float)
        self.assertEqual(wb["CPI"]["A5"].value, "2026-09")
        self.assertIn("VDVGold", wb["Vàng SJC"].tables)
        self.assertIn("VDVCpi", wb["CPI"].tables)
        self.assertEqual(wb.defined_names["VDV_PERIOD"].attr_text, "'Bắt đầu'!$C$7")

    def test_workbook_embeds_standard_office_addin_parts(self):
        response = self._response()
        with zipfile.ZipFile(io.BytesIO(response.body)) as package:
            self.assertIsNone(package.testzip())
            names = set(package.namelist())
            self.assertIn("xl/webextensions/taskpanes.xml", names)
            self.assertIn("xl/webextensions/webextension1.xml", names)
            self.assertIn("xl/webextensions/_rels/taskpanes.xml.rels", names)
            extension = package.read("xl/webextensions/webextension1.xml").decode()
            self.assertIn(excel_export.ADDIN_ID, extension)
            self.assertIn("Office.AutoShowTaskpaneWithDocument", extension)

    def test_workbook_never_contains_an_api_key(self):
        response = self._response()
        self.assertNotIn(b"vdv_secret", response.body)
        self.assertNotIn(b"api_key=", response.body)

    def test_refresh_payload_is_one_json_bundle_for_all_tables(self):
        data = _datasets()
        with patch.object(excel_export, "_collect_datasets", AsyncMock(return_value=data)):
            payload = asyncio.run(excel_export.excel_refresh_data(_Request(), "1y"))
        encoded = jsonable_encoder(payload)
        self.assertTrue(encoded["success"])
        self.assertEqual(encoded["count"], 9)
        self.assertEqual(encoded["source"], "Viet Dataverse")
        self.assertEqual(encoded["data"][0]["rows"][0][0], "2026-09-23")
        self.assertEqual({item["table"] for item in encoded["data"]}, {
            "VDVGold", "VDVSilver", "VDVInterbank", "VDVFx", "VDVDeposit",
            "VDVGlobal", "VDVCpi", "VDVGdp", "VDVTrade",
        })


if __name__ == "__main__":
    unittest.main()
